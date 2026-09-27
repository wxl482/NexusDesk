# -*- coding: utf-8 -*-
"""
生成《AI Agent 开发实战指南》PPT
结合真实工程（FastAPI + LangChain + LangGraph）定制
"""
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn
import copy

# ---------------- 配色 ----------------
INK     = RGBColor(0x0F, 0x17, 0x2A)
SUB     = RGBColor(0x47, 0x55, 0x69)
MUTED   = RGBColor(0x94, 0xA3, 0xB8)
BLUE    = RGBColor(0x25, 0x63, 0xEB)
DEEP    = RGBColor(0x1B, 0x2A, 0x5E)
DEEP2   = RGBColor(0x0F, 0x17, 0x2A)
CYAN    = RGBColor(0x06, 0xB6, 0xD4)
PURPLE  = RGBColor(0x7C, 0x3A, 0xED)
ORANGE  = RGBColor(0xF5, 0x9E, 0x0B)
GREEN   = RGBColor(0x10, 0xB9, 0x81)
LIGHT   = RGBColor(0xEF, 0xF3, 0xF9)
CARD    = RGBColor(0xF8, 0xFA, 0xFC)
WHITE   = RGBColor(0xFF, 0xFF, 0xFF)
BORDER  = RGBColor(0xE2, 0xE8, 0xF0)
CODEBG  = RGBColor(0x0D, 0x1B, 0x2A)

FONT = 'Microsoft YaHei'
MONO = 'Consolas'

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]
SW, SH = 13.333, 7.5

# ---------------- 工具函数 ----------------
def set_font(run, name=FONT, size=18, bold=False, color=INK, italic=False):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = color
    run.font.name = name
    rPr = run._r.get_or_add_rPr()
    for tag in ('a:ea', 'a:cs'):
        el = rPr.find(qn(tag))
        if el is None:
            el = rPr.makeelement(qn(tag), {})
            rPr.append(el)
        el.set('typeface', name)

def add_text(slide, l, t, w, h, items, anchor=MSO_ANCHOR.TOP, wrap=True):
    tb = slide.shapes.add_textbox(Inches(l), Inches(t), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = wrap
    tf.vertical_anchor = anchor
    tf.margin_left = 0; tf.margin_right = 0
    tf.margin_top = 0; tf.margin_bottom = 0
    first = True
    for it in items:
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.alignment = it.get('align', PP_ALIGN.LEFT)
        p.line_spacing = it.get('ls', 1.12)
        p.space_after = Pt(it.get('sa', 6))
        p.space_before = Pt(it.get('sb', 0))
        run = p.add_run()
        run.text = it.get('text', '')
        set_font(run, it.get('font', FONT), it.get('size', 18),
                 it.get('bold', False), it.get('color', INK))
    return tb

def add_rect(slide, l, t, w, h, fill=None, line=None, shape=MSO_SHAPE.ROUNDED_RECTANGLE,
             line_w=1.0, radius=0.08):
    sp = slide.shapes.add_shape(shape, Inches(l), Inches(t), Inches(w), Inches(h))
    if fill is None:
        sp.fill.background()
    else:
        sp.fill.solid(); sp.fill.fore_color.rgb = fill
    if line is None:
        sp.line.fill.background()
    else:
        sp.line.color.rgb = line; sp.line.width = Pt(line_w)
    sp.shadow.inherit = False
    if shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        try:
            sp.adjustments[0] = radius
        except Exception:
            pass
    return sp

def set_gradient(shape, c1, c2, angle=45):
    shape.fill.gradient()
    shape.fill.gradient_angle = angle
    stops = shape.fill.gradient_stops
    stops[0].color.rgb = c1; stops[0].position = 0.0
    stops[1].color.rgb = c2; stops[1].position = 1.0

def header(slide, no, title, en=''):
    add_rect(slide, 0.0, 0.0, 0.14, SH, fill=BLUE, shape=MSO_SHAPE.RECTANGLE)
    add_text(slide, 0.75, 0.42, 11.5, 0.3,
             [{'text': en, 'size': 12, 'bold': True, 'color': BLUE}])
    add_text(slide, 0.75, 0.68, 11.5, 0.7,
             [{'text': title, 'size': 27, 'bold': True, 'color': INK}])
    add_rect(slide, 0.78, 1.38, 1.0, 0.06, fill=BLUE, shape=MSO_SHAPE.RECTANGLE)

def footer(slide, no):
    add_text(slide, 11.6, 7.02, 1.3, 0.3,
             [{'text': f'{no:02d}', 'size': 11, 'bold': True, 'color': MUTED,
               'align': PP_ALIGN.RIGHT}])
    add_text(slide, 0.75, 7.02, 6.0, 0.3,
             [{'text': 'AI Agent 开发实战指南', 'size': 10, 'color': MUTED}])

def chip(slide, l, t, text, color, w=None, size=12):
    w = w or (0.16 + len(text) * 0.145)
    sp = add_rect(slide, l, t, w, 0.36, fill=color, radius=0.5)
    tf = sp.text_frame; tf.word_wrap = False
    tf.margin_left = Inches(0.1); tf.margin_right = Inches(0.1)
    tf.margin_top = 0; tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
    r = p.add_run(); r.text = text
    set_font(r, FONT, size, True, WHITE)
    return sp

def number_badge(slide, l, t, num, color, d=0.5, size=18):
    sp = add_rect(slide, l, t, d, d, fill=color, shape=MSO_SHAPE.OVAL)
    tf = sp.text_frame; tf.margin_left = 0; tf.margin_right = 0
    tf.margin_top = 0; tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
    r = p.add_run(); r.text = str(num)
    set_font(r, FONT, size, True, WHITE)
    return sp

# ---------------- 1. 封面 ----------------
def slide_cover():
    s = prs.slides.add_slide(BLANK)
    bg = add_rect(s, 0, 0, SW, SH, fill=DEEP2, shape=MSO_SHAPE.RECTANGLE)
    set_gradient(bg, DEEP2, DEEP, angle=45)
    # 装饰圆
    c1 = add_rect(s, 9.6, -1.4, 5.2, 5.2, fill=BLUE, shape=MSO_SHAPE.OVAL)
    c1.fill.fore_color.rgb = RGBColor(0x1E, 0x3A, 0x8A)
    c2 = add_rect(s, 11.2, 3.6, 3.4, 3.4, fill=CYAN, shape=MSO_SHAPE.OVAL)
    c2.fill.fore_color.rgb = RGBColor(0x0E, 0x4A, 0x63)
    c3 = add_rect(s, 10.6, 1.0, 1.4, 1.4, fill=CYAN, shape=MSO_SHAPE.OVAL)
    # 标签
    chip(s, 0.9, 1.5, '  AI · AGENT · LLM  ', CYAN, w=2.7, size=13)
    add_text(s, 0.9, 2.15, 10.5, 1.6,
             [{'text': 'AI Agent', 'size': 62, 'bold': True, 'color': WHITE, 'ls': 1.0}])
    add_text(s, 0.9, 3.35, 11.0, 1.6,
             [{'text': '开发实战指南', 'size': 54, 'bold': True, 'color': CYAN, 'ls': 1.0}])
    add_rect(s, 0.95, 4.35, 2.2, 0.07, fill=ORANGE, shape=MSO_SHAPE.RECTANGLE)
    add_text(s, 0.9, 4.6, 11.0, 1.0,
             [{'text': '从核心原理到 LangGraph 工程落地', 'size': 22, 'color': RGBColor(0xCB,0xD5,0xE1)}])
    add_text(s, 0.9, 6.3, 11.0, 0.9, [
        {'text': '关键词：LLM 推理 · 工具调用 · RAG 记忆 · 多智能体协作 · 状态图工作流',
         'size': 13, 'color': RGBColor(0x8A,0x9B,0xB0)},
    ])

# ---------------- 2. 目录 ----------------
def slide_toc():
    s = prs.slides.add_slide(BLANK)
    header(s, 2, '目录 Contents', 'AGENDA')
    items = [
        ('01', '什么是 AI Agent', '从 LLM 到自主智能体的跃迁', BLUE),
        ('02', 'Agent 核心架构', '感知 · 规划 · 记忆 · 行动', CYAN),
        ('03', '工作范式与推理', 'ReAct / Plan-Execute / Reflection', PURPLE),
        ('04', '工具调用与记忆机制', 'Function Calling · RAG · Checkpointer', ORANGE),
        ('05', '多智能体协作', 'Orchestrator · 层级 · 群聊', GREEN),
        ('06', '框架选型与工程落地', 'LangGraph 实战架构解析', BLUE),
        ('07', '最佳实践与未来趋势', 'MCP · A2A · 可观测性', CYAN),
    ]
    x0, y0 = 0.9, 1.75
    col_w, row_h = 5.9, 0.72
    for i, (no, title, desc, color) in enumerate(items):
        col = i % 2; row = i // 2
        l = x0 + col * col_w
        t = y0 + row * row_h
        number_badge(s, l, t + 0.02, no, color, d=0.5, size=15)
        add_text(s, l + 0.7, t, 4.9, 0.75, [
            {'text': title, 'size': 16, 'bold': True, 'color': INK, 'sa': 1},
            {'text': desc, 'size': 11.5, 'color': SUB},
        ])

# ---------------- 3. 什么是 Agent ----------------
def slide_what_is():
    s = prs.slides.add_slide(BLANK)
    header(s, 3, '什么是 AI Agent', 'CONCEPT')
    sp = add_rect(s, 0.75, 1.65, 11.85, 1.15, fill=LIGHT, line=BORDER)
    add_text(s, 1.1, 1.85, 11.2, 0.5, [
        {'text': 'AI Agent（智能体）', 'size': 17, 'bold': True, 'color': DEEP, 'sa': 2},
        {'text': '是以大语言模型（LLM）为「大脑」，能够自主感知环境、规划任务、'
                 '调用工具并执行行动，最终达成目标的系统。', 'size': 14, 'color': SUB},
    ])
    # 公式
    add_text(s, 0.75, 3.05, 11.85, 0.5,
             [{'text': '核心公式', 'size': 14, 'bold': True, 'color': INK}])
    parts = [('Agent', BLUE), ('=', MUTED), ('LLM', PURPLE), ('+', MUTED),
             ('规划 Planning', CYAN), ('+', MUTED), ('记忆 Memory', ORANGE),
             ('+', MUTED), ('工具 Tools', GREEN)]
    x = 0.9; y = 3.55
    for text, color in parts:
        w = 0.28 + len(text) * 0.185
        if text in ('=', '+'):
            add_text(s, x, y, 0.4, 0.5, [{'text': text, 'size': 20, 'bold': True,
                    'color': MUTED, 'align': PP_ALIGN.CENTER}])
            x += 0.42
        else:
            add_rect(s, x, y, w, 0.55, fill=color, radius=0.25)
            add_text(s, x, y + 0.06, w, 0.45, [{'text': text, 'size': 15, 'bold': True,
                     'color': WHITE, 'align': PP_ALIGN.CENTER}])
            x += w + 0.12
    # 对比表
    data = [
        ('对比维度', '传统 LLM 对话', 'AI Agent'),
        ('核心能力', '生成文本 / 回答问题', '自主决策 + 执行行动'),
        ('交互方式', '单轮 / 多轮问答', '目标驱动，多步自主循环'),
        ('是否用工具', '否', '是（搜索/代码/API 等）'),
        ('典型形态', 'ChatBot 聊天机器人', '自动驾驶式的任务执行体'),
    ]
    l, t, w, h = 0.75, 4.35, 11.85, 2.2
    gf = s.shapes.add_table(len(data), 3, Inches(l), Inches(t), Inches(w), Inches(h))
    tbl = gf.table
    tbl.columns[0].width = Inches(2.2)
    tbl.columns[1].width = Inches(4.5)
    tbl.columns[2].width = Inches(5.15)
    for r, row in enumerate(data):
        for c, val in enumerate(row):
            cell = tbl.cell(r, c)
            cell.margin_left = Inches(0.18); cell.margin_right = Inches(0.1)
            cell.margin_top = Inches(0.05); cell.margin_bottom = Inches(0.05)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            tf = cell.text_frame; tf.word_wrap = True
            p = tf.paragraphs[0]; p.alignment = PP_ALIGN.LEFT
            run = p.add_run(); run.text = val
            if r == 0:
                cell.fill.solid(); cell.fill.fore_color.rgb = DEEP
                set_font(run, FONT, 13.5, True, WHITE)
            else:
                cell.fill.solid()
                cell.fill.fore_color.rgb = WHITE if r % 2 else LIGHT
                set_font(run, FONT, 12.5, c == 0, INK if c == 0 else SUB)
    footer(s, 3)

# ---------------- 4. 核心架构 ----------------
def slide_architecture():
    s = prs.slides.add_slide(BLANK)
    header(s, 4, 'Agent 核心架构', 'ARCHITECTURE')
    # 中心大脑
    brain = add_rect(s, 5.16, 3.15, 3.0, 1.3, fill=DEEP, radius=0.15)
    add_text(s, 5.16, 3.35, 3.0, 0.95, [
        {'text': 'LLM 核心大脑', 'size': 17, 'bold': True, 'color': WHITE, 'align': PP_ALIGN.CENTER, 'sa': 4},
        {'text': '推理 · 决策 · 语言理解', 'size': 11, 'color': RGBColor(0xBE,0xD0,0xF0), 'align': PP_ALIGN.CENTER},
    ])
    pillars = [
        (0.75, 1.75, '感知 Perception', '接收用户指令、环境状态\n与多模态输入', CYAN),
        (9.85, 1.75, '记忆 Memory', '短期对话上下文 +\n长期知识库 (RAG)', ORANGE),
        (0.75, 4.85, '规划 Planning', '任务拆解、路径推理、\n步骤编排与自我反思', PURPLE),
        (9.85, 4.85, '行动 Action', '调用工具、执行代码、\n操作文件与外部 API', GREEN),
    ]
    for l, t, title, desc, color in pillars:
        add_rect(s, l, t, 3.2, 1.7, fill=CARD, line=BORDER)
        add_rect(s, l, t, 0.12, 1.7, fill=color, shape=MSO_SHAPE.RECTANGLE)
        add_text(s, l + 0.35, t + 0.25, 2.75, 1.2, [
            {'text': title, 'size': 15, 'bold': True, 'color': color, 'sa': 5},
            {'text': desc, 'size': 12, 'color': SUB, 'ls': 1.25},
        ])
    # 连接箭头（用细线示意）
    for l, t in [(3.95, 2.55), (8.85, 2.55), (3.95, 5.45), (8.85, 5.45)]:
        add_rect(s, l, t, 1.2, 0.04, fill=BORDER, shape=MSO_SHAPE.RECTANGLE)
    footer(s, 4)

# ---------------- 5. 工作范式 ----------------
def slide_paradigms():
    s = prs.slides.add_slide(BLANK)
    header(s, 5, 'Agent 工作范式与推理', 'PARADIGMS')
    cards = [
        ('ReAct', 'Reasoning + Acting', '思考与行动交替进行，形成循环：\nThought → Action → Observation，\n边推理边调用工具，动态修正。', BLUE, '主流首选'),
        ('Plan-and-Execute', '先规划后执行', '先由规划器生成完整步骤清单，\n再逐步执行；适合流程明确、\n可预分解的复杂任务。', CYAN, '长链路任务'),
        ('Reflection', '自我反思迭代', 'Agent 审视自身输出，\n主动批评与修正，\n多轮打磨提升结果质量。', PURPLE, '质量优先'),
    ]
    x = 0.75; y = 1.75; w = 3.78; h = 4.7; gap = 0.27
    for i, (title, en, desc, color, tag) in enumerate(cards):
        l = x + i * (w + gap)
        add_rect(s, l, y, w, h, fill=WHITE, line=BORDER)
        add_rect(s, l, y, w, 0.14, fill=color, shape=MSO_SHAPE.RECTANGLE)
        number_badge(s, l + 0.35, y + 0.45, f'0{i+1}', color, d=0.62, size=18)
        chip(s, l + 2.05, y + 0.55, tag, color, size=11)
        add_text(s, l + 0.35, y + 1.35, w - 0.7, 1.0, [
            {'text': title, 'size': 21, 'bold': True, 'color': INK, 'sa': 2},
            {'text': en, 'size': 12, 'color': color, 'bold': True},
        ])
        add_text(s, l + 0.35, y + 2.6, w - 0.7, 1.9, [
            {'text': desc, 'size': 13, 'color': SUB, 'ls': 1.4},
        ])
    footer(s, 5)

# ---------------- 6. 工具调用 ----------------
def slide_tools():
    s = prs.slides.add_slide(BLANK)
    header(s, 6, '工具调用 · Function Calling', 'TOOL USE')
    add_text(s, 0.75, 1.62, 11.85, 0.4, [
        {'text': '工具是 Agent 的「手和脚」——让 LLM 能够真正与外部世界交互。', 'size': 13.5, 'color': SUB}])
    # 左：三要素
    add_rect(s, 0.75, 2.15, 5.7, 4.3, fill=CARD, line=BORDER)
    add_text(s, 1.05, 2.4, 5.1, 0.5, [{'text': '工具定义三要素', 'size': 16, 'bold': True, 'color': DEEP}])
    elems = [('名称 Name', '唯一标识，如 execute_python_code', BLUE),
             ('描述 Description', '自然语言说明用途，直接影响模型选择', CYAN),
             ('参数 Schema', '结构化参数定义，约束输入格式', PURPLE)]
    for i, (t1, t2, color) in enumerate(elems):
        yy = 3.0 + i * 1.05
        add_rect(s, 1.05, yy, 0.1, 0.85, fill=color, shape=MSO_SHAPE.RECTANGLE)
        add_text(s, 1.3, yy - 0.02, 4.9, 0.9, [
            {'text': t1, 'size': 14, 'bold': True, 'color': color, 'sa': 2},
            {'text': t2, 'size': 11.5, 'color': SUB, 'ls': 1.2},
        ])
    # 右：流程
    add_rect(s, 6.9, 2.15, 5.7, 4.3, fill=WHITE, line=BORDER)
    add_text(s, 7.2, 2.4, 5.1, 0.5, [{'text': '调用闭环流程', 'size': 16, 'bold': True, 'color': DEEP}])
    flow = [('1', '模型决定调用工具', '输出 tool_calls 结构化请求', BLUE),
            ('2', '框架执行工具', 'ToolNode 运行并返回 Observation', CYAN),
            ('3', '结果回填模型', '作为 ToolMessage 注入上下文', PURPLE),
            ('4', '模型生成最终答复', '或继续下一轮工具调用', GREEN)]
    for i, (n, t1, t2, color) in enumerate(flow):
        yy = 2.95 + i * 0.85
        number_badge(s, 7.2, yy, n, color, d=0.42, size=13)
        add_text(s, 7.8, yy - 0.03, 4.6, 0.8, [
            {'text': t1, 'size': 13, 'bold': True, 'color': INK, 'sa': 1},
            {'text': t2, 'size': 11, 'color': SUB},
        ])
    # 底部代码
    add_rect(s, 0.75, 6.55, 11.85, 0.55, fill=CODEBG)
    add_text(s, 1.0, 6.68, 11.5, 0.4, [
        {'text': 'llm_with_tools = base_llm.bind_tools(DEFAULT_TOOLS)   # LangChain 一行完成工具绑定',
         'size': 11.5, 'color': RGBColor(0x7D,0xE0,0xC0), 'font': MONO}])
    footer(s, 6)

# ---------------- 7. 记忆机制 ----------------
def slide_memory():
    s = prs.slides.add_slide(BLANK)
    header(s, 7, '记忆机制 · Memory & RAG', 'MEMORY')
    cards = [
        ('短期记忆', 'Short-term', ['完整对话 messages 列表',
                                    '滑动窗口 / Token 截断',
                                    '保障多轮上下文连贯'], BLUE),
        ('长期记忆', 'Long-term / RAG', ['向量数据库存储知识',
                                        '语义检索相关片段',
                                        '突破上下文长度限制'], CYAN),
        ('状态持久化', 'Checkpointer', ['MemorySaver / SqliteSaver',
                                      '按 thread_id 保存会话',
                                      '断点续跑、跨会话记忆'], PURPLE),
    ]
    x = 0.75; y = 1.75; w = 3.78; h = 3.5; gap = 0.27
    for i, (title, en, pts, color) in enumerate(cards):
        l = x + i * (w + gap)
        add_rect(s, l, y, w, h, fill=WHITE, line=BORDER)
        add_rect(s, l, y, w, 1.05, fill=color, radius=0.08)
        add_text(s, l + 0.3, y + 0.16, w - 0.6, 0.8, [
            {'text': title, 'size': 17, 'bold': True, 'color': WHITE, 'sa': 2},
            {'text': en, 'size': 11, 'color': RGBColor(0xE6,0xEE,0xFB)}])
        tb = add_text(s, l + 0.32, y + 1.3, w - 0.64, 2.0, [])
        tf = tb.text_frame
        first = True
        for pt in pts:
            p = tf.paragraphs[0] if first else tf.add_paragraph()
            first = False
            p.line_spacing = 1.3; p.space_after = Pt(8)
            r = p.add_run(); r.text = '•  ' + pt
            set_font(r, FONT, 12.5, False, SUB)
    add_rect(s, 0.75, 5.55, 11.85, 1.15, fill=LIGHT, line=BORDER)
    add_text(s, 1.1, 5.75, 11.2, 0.85, [
        {'text': '设计要点', 'size': 13, 'bold': True, 'color': DEEP, 'sa': 3},
        {'text': '记忆不是「存得越多越好」——需通过摘要压缩、相关性检索与检索增强（RAG），'
                 '在有限上下文窗口内保留最有效信息。', 'size': 12.5, 'color': SUB}])
    footer(s, 7)

# ---------------- 8. 多智能体 ----------------
def slide_multiagent():
    s = prs.slides.add_slide(BLANK)
    header(s, 8, '多智能体协作', 'MULTI-AGENT')
    add_text(s, 0.75, 1.62, 11.85, 0.4, [
        {'text': '将复杂任务拆解给多个专业化 Agent 协同完成，分工明确、各司其职。',
         'size': 13.5, 'color': SUB}])
    modes = [
        ('Orchestrator 调度', '总调度官拆解任务并分派给专业 Agent，\n汇总结果统一交付。', BLUE),
        ('Hierarchical 层级式', '上下级 Agent 树状结构，\n逐层下发与上报，适合超复杂任务。', CYAN),
        ('Group Chat 群聊', '多个 Agent 在同一会话中\n讨论辩论，达成共识。', PURPLE),
        ('Pipeline 流水线', '按固定顺序依次处理，\n前一步输出作为后一步输入。', ORANGE),
    ]
    x = 0.75; y = 2.3; w = 5.78; h = 1.8; gapx = 0.3; gapy = 0.3
    for i, (title, desc, color) in enumerate(modes):
        col = i % 2; row = i // 2
        l = x + col * (w + gapx); t = y + row * (h + gapy)
        add_rect(s, l, t, w, h, fill=CARD, line=BORDER)
        add_rect(s, l, t, 0.5, h, fill=color, shape=MSO_SHAPE.RECTANGLE)
        add_text(s, l + 0.75, t + 0.28, w - 1.0, 1.3, [
            {'text': title, 'size': 15.5, 'bold': True, 'color': color, 'sa': 5},
            {'text': desc, 'size': 12.5, 'color': SUB, 'ls': 1.3}])
    add_rect(s, 0.75, 6.35, 11.85, 0.55, fill=LIGHT, line=BORDER)
    add_text(s, 1.05, 6.46, 11.4, 0.4, [
        {'text': '价值：专业化分工 · 并行提速 · 相互校验提升鲁棒性 · 突破单 Agent 能力上限',
         'size': 12.5, 'bold': True, 'color': DEEP}])
    footer(s, 8)

# ---------------- 9. 框架选型 ----------------
def slide_frameworks():
    s = prs.slides.add_slide(BLANK)
    header(s, 9, '主流开发框架选型', 'FRAMEWORKS')
    data = [
        ('框架', '定位', '特点', '适用场景'),
        ('LangChain', 'LLM 应用开发基石', '组件丰富、生态最广、LCEL 编排', '通用 LLM 应用'),
        ('LangGraph', 'Agent 状态图编排', '图状工作流、循环/条件边、Checkpoint', '复杂可控 Agent ★'),
        ('AutoGen', '多智能体对话', '微软出品，Agent 群聊协作', '多 Agent 研究'),
        ('CrewAI', '角色化多智能体', '角色/任务抽象直观，上手快', '业务流程自动化'),
    ]
    l, t, w, h = 0.75, 1.7, 11.85, 3.9
    gf = s.shapes.add_table(len(data), 4, Inches(l), Inches(t), Inches(w), Inches(h))
    tbl = gf.table
    for i, wd in enumerate([2.4, 2.9, 4.35, 2.2]):
        tbl.columns[i].width = Inches(wd)
    for r, row in enumerate(data):
        for c, val in enumerate(row):
            cell = tbl.cell(r, c)
            cell.margin_left = Inches(0.16); cell.margin_right = Inches(0.08)
            cell.margin_top = Inches(0.06); cell.margin_bottom = Inches(0.06)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            tf = cell.text_frame; tf.word_wrap = True
            p = tf.paragraphs[0]; p.alignment = PP_ALIGN.LEFT
            run = p.add_run(); run.text = val
            if r == 0:
                cell.fill.solid(); cell.fill.fore_color.rgb = DEEP
                set_font(run, FONT, 13.5, True, WHITE)
            else:
                cell.fill.solid()
                cell.fill.fore_color.rgb = RGBColor(0xE8,0xF0,0xFE) if r == 2 else (WHITE if r % 2 else LIGHT)
                set_font(run, FONT, 12.5, c == 0, DEEP if c == 0 else SUB)
    add_rect(s, 0.75, 5.85, 11.85, 0.95, fill=LIGHT, line=BORDER)
    add_text(s, 1.1, 6.02, 11.2, 0.7, [
        {'text': '选型建议', 'size': 13, 'bold': True, 'color': DEEP, 'sa': 3},
        {'text': '需要流程可控、可循环、可持久化的生产级 Agent → 首选 LangGraph；'
                 '多智能体协作场景可叠加 AutoGen / CrewAI。', 'size': 12.5, 'color': SUB}])
    footer(s, 9)

# ---------------- 10. 实战架构 ----------------
def slide_case():
    s = prs.slides.add_slide(BLANK)
    header(s, 10, '实战架构解析：本项目分层设计', 'CASE STUDY')
    layers = [
        ('API 层  (FastAPI)', 'chat · rag · tools · health  —  提供 SSE 流式接口', BLUE),
        ('Agent 层  (LangGraph)', 'StateGraph 状态图  ·  agent 节点  ·  tools 节点  ·  条件路由', PURPLE),
        ('Chains 层  (LCEL)', 'planner_chain 规划  ·  rag_chain 检索  ·  summary_chain 摘要', CYAN),
        ('Tools 层', 'terminal · file_ops · search · python_exec · query_knowledge_base', GREEN),
        ('RAG / LLM 层', 'ChromaDB 向量引擎  ·  LLMFactory 模型工厂（可切换）', ORANGE),
    ]
    y = 1.75; h = 0.92; gap = 0.13
    for i, (title, desc, color) in enumerate(layers):
        t = y + i * (h + gap)
        add_rect(s, 0.75, t, 11.85, h, fill=CARD, line=BORDER)
        add_rect(s, 0.75, t, 0.16, h, fill=color, shape=MSO_SHAPE.RECTANGLE)
        add_text(s, 1.15, t + 0.13, 11.2, h - 0.2, [
            {'text': title, 'size': 14.5, 'bold': True, 'color': color, 'sa': 2},
            {'text': desc, 'size': 11.5, 'color': SUB}])
    footer(s, 10)

# ---------------- 11. LangGraph 工作流 ----------------
def slide_langgraph():
    s = prs.slides.add_slide(BLANK)
    header(s, 11, 'LangGraph 状态图工作流', 'WORKFLOW')
    # 流程图
    def node(l, t, w, h, text, color, sub=None):
        add_rect(s, l, t, w, h, fill=color, radius=0.2)
        items = [{'text': text, 'size': 13, 'bold': True, 'color': WHITE,
                  'align': PP_ALIGN.CENTER, 'sa': 1}]
        if sub:
            items.append({'text': sub, 'size': 10, 'color': RGBColor(0xDD,0xE7,0xF5),
                          'align': PP_ALIGN.CENTER})
        add_text(s, l, t + 0.16, w, h, items, anchor=MSO_ANCHOR.MIDDLE)
    node(0.9, 3.35, 2.4, 0.95, 'START', INK)
    node(3.7, 3.35, 2.6, 0.95, 'agent 节点', PURPLE, 'LLM 推理')
    node(6.9, 1.9, 2.6, 0.95, 'tools 节点', GREEN, 'ToolNode 执行')
    node(10.0, 3.35, 2.4, 0.95, 'END', INK)
    add_text(s, 6.9, 3.05, 2.6, 0.4, [{'text': '有工具调用 →', 'size': 10.5, 'color': GREEN,
             'align': PP_ALIGN.CENTER, 'bold': True}])
    add_text(s, 6.9, 4.35, 2.6, 0.4, [{'text': '← 结果回环', 'size': 10.5, 'color': GREEN,
             'align': PP_ALIGN.CENTER, 'bold': True}])
    for l, t, w, h, c in [(3.3, 3.74, 0.4, 0.05, MUTED), (6.3, 3.74, 0.6, 0.05, MUTED),
                          (9.5, 3.74, 0.5, 0.05, MUTED)]:
        add_rect(s, l, t, w, h, fill=c, shape=MSO_SHAPE.RECTANGLE)
    # 特点
    feats = [('StateGraph', '状态图统一编排，节点 + 边'), ('tools_condition', '自动判断是否调用工具'),
             ('Checkpointer', 'MemorySaver 持久化会话'), ('循环能力', 'Agent ↔ Tools 闭环迭代')]
    for i, (a, b) in enumerate(feats):
        l = 0.9 + i * 3.05
        add_rect(s, l, 5.6, 2.85, 1.15, fill=LIGHT, line=BORDER)
        add_text(s, l + 0.25, 5.78, 2.4, 0.9, [
            {'text': a, 'size': 13, 'bold': True, 'color': DEEP, 'sa': 3},
            {'text': b, 'size': 10.5, 'color': SUB, 'ls': 1.15}])
    footer(s, 11)

# ---------------- 12. 核心代码 ----------------
def slide_code():
    s = prs.slides.add_slide(BLANK)
    header(s, 12, '核心代码：组装 Agent 工作流', 'CODE')
    add_rect(s, 0.75, 1.7, 11.85, 4.95, fill=CODEBG, radius=0.04)
    lines = [
        ('from langgraph.graph import StateGraph, START, END', RGBColor(0x7D,0xE0,0xC0)),
        ('from langgraph.prebuilt import ToolNode, tools_condition', RGBColor(0x7D,0xE0,0xC0)),
        ('', INK),
        ('workflow = StateGraph(AgentState)                # 1. 定义状态图', RGBColor(0xE2,0xE8,0xF0)),
        ('workflow.add_node("agent", agent_node)           # 2. 推理节点', RGBColor(0xE2,0xE8,0xF0)),
        ('workflow.add_node("tools", ToolNode(DEFAULT_TOOLS))  # 3. 工具节点', RGBColor(0xE2,0xE8,0xF0)),
        ('', INK),
        ('workflow.add_edge(START, "agent")                # 4. 入口', RGBColor(0x9C,0xB4,0xF0)),
        ('workflow.add_conditional_edges(', RGBColor(0x9C,0xB4,0xF0)),
        ('    "agent", tools_condition,                    # 5. 条件路由', RGBColor(0x9C,0xB4,0xF0)),
        ('    {"tools": "tools", END: END})', RGBColor(0x9C,0xB4,0xF0)),
        ('workflow.add_edge("tools", "agent")              # 6. 回环闭环', RGBColor(0x9C,0xB4,0xF0)),
        ('', INK),
        ('app = workflow.compile(checkpointer=memory_checkpointer)  # 7. 编译+持久化', RGBColor(0xF5,0xB0,0x6B)),
    ]
    items = []
    for txt, col in lines:
        items.append({'text': txt if txt else ' ', 'font': MONO, 'size': 13.5,
                      'color': col, 'ls': 1.25, 'sa': 0})
    add_text(s, 1.15, 1.95, 11.1, 4.5, items)
    footer(s, 12)

# ---------------- 13. 最佳实践 ----------------
def slide_best():
    s = prs.slides.add_slide(BLANK)
    header(s, 13, '开发最佳实践', 'BEST PRACTICES')
    items = [
        ('提示词与工具描述', '工具描述要精准清晰，直接决定模型选择正确率', BLUE),
        ('状态设计', '合理定义 State 结构，善用 Annotated 归约器管理消息', PURPLE),
        ('错误处理', '为工具加超时、重试与降级，避免 Agent 卡死', ORANGE),
        ('可观测性', '接入 LangSmith 等链路追踪，记录每一步 Thought/Action', CYAN),
        ('成本与延迟', '控制上下文长度、缓存结果、按需选择小模型', GREEN),
        ('安全边界', '工具沙箱化、权限最小化，人工审核高危操作', RGBColor(0xEF,0x44,0x44)),
    ]
    x = 0.75; y = 1.7; w = 5.78; h = 1.28; gapx = 0.3; gapy = 0.18
    for i, (title, desc, color) in enumerate(items):
        col = i % 2; row = i // 2
        l = x + col * (w + gapx); t = y + row * (h + gapy)
        add_rect(s, l, t, w, h, fill=WHITE, line=BORDER)
        add_rect(s, l, t, 0.5, h, fill=color, radius=0.0)
        number_badge(s, l + 0.72, t + 0.37, f'{i+1}', color, d=0.52, size=15)
        add_text(s, l + 1.4, t + 0.22, w - 1.6, 0.95, [
            {'text': title, 'size': 14.5, 'bold': True, 'color': INK, 'sa': 3},
            {'text': desc, 'size': 11.5, 'color': SUB, 'ls': 1.2}])
    footer(s, 13)

# ---------------- 14. 未来趋势 ----------------
def slide_future():
    s = prs.slides.add_slide(BLANK)
    header(s, 14, '挑战与未来趋势', 'FUTURE')
    left = [
        ('MCP 协议', 'Model Context Protocol：统一工具/上下文接入标准', CYAN),
        ('A2A 协作', 'Agent-to-Agent：智能体之间的通信与协作协议', BLUE),
        ('多模态 Agent', '融合文本、图像、语音、视频的感知与行动', PURPLE),
        ('持续学习记忆', '从交互中沉淀长期经验，实现自我进化', ORANGE),
    ]
    add_text(s, 0.9, 1.65, 5.5, 0.5, [{'text': '▍ 发展趋势', 'size': 16, 'bold': True, 'color': DEEP}])
    for i, (t1, t2, color) in enumerate(left):
        t = 2.25 + i * 1.06
        add_rect(s, 0.9, t, 5.6, 0.92, fill=CARD, line=BORDER)
        add_rect(s, 0.9, t, 0.1, 0.92, fill=color, shape=MSO_SHAPE.RECTANGLE)
        add_text(s, 1.2, t + 0.13, 5.1, 0.7, [
            {'text': t1, 'size': 14, 'bold': True, 'color': color, 'sa': 2},
            {'text': t2, 'size': 11.5, 'color': SUB}])
    right = [
        ('可靠性', '幻觉、错误累积与长链路稳定性'),
        ('评估难题', '缺乏统一、客观的 Agent 能力评测标准'),
        ('成本控制', '多步推理与调用带来的 Token 开销'),
        ('安全合规', '越权操作、数据隐私与可控性'),
    ]
    add_text(s, 7.0, 1.65, 5.5, 0.5, [{'text': '▍ 现存挑战', 'size': 16, 'bold': True, 'color': DEEP}])
    for i, (t1, t2) in enumerate(right):
        t = 2.25 + i * 1.06
        add_rect(s, 7.0, t, 5.6, 0.92, fill=WHITE, line=BORDER)
        number_badge(s, 7.25, t + 0.22, f'{i+1}', MUTED, d=0.48, size=14)
        add_text(s, 7.95, t + 0.13, 4.5, 0.7, [
            {'text': t1, 'size': 14, 'bold': True, 'color': INK, 'sa': 2},
            {'text': t2, 'size': 11.5, 'color': SUB}])
    footer(s, 14)

# ---------------- 15. 结尾 ----------------
def slide_thanks():
    s = prs.slides.add_slide(BLANK)
    bg = add_rect(s, 0, 0, SW, SH, fill=DEEP2, shape=MSO_SHAPE.RECTANGLE)
    set_gradient(bg, DEEP2, DEEP, angle=45)
    c2 = add_rect(s, 10.8, -1.2, 4.0, 4.0, fill=RGBColor(0x1E,0x3A,0x8A), shape=MSO_SHAPE.OVAL)
    c3 = add_rect(s, 11.6, 4.2, 3.0, 3.0, fill=RGBColor(0x0E,0x4A,0x63), shape=MSO_SHAPE.OVAL)
    add_text(s, 0.9, 2.7, 11.5, 1.5, [
        {'text': 'Thank You', 'size': 60, 'bold': True, 'color': WHITE, 'align': PP_ALIGN.CENTER}])
    add_rect(s, 5.67, 4.35, 2.0, 0.06, fill=ORANGE, shape=MSO_SHAPE.RECTANGLE)
    add_text(s, 0.9, 4.65, 11.5, 0.8, [
        {'text': '让每一个目标，都能被自主地拆解与达成', 'size': 18,
         'color': RGBColor(0xCB,0xD5,0xE1), 'align': PP_ALIGN.CENTER}])
    add_text(s, 0.9, 6.5, 11.5, 0.5, [
        {'text': 'Agent 开发实战指南  ·  LangChain + LangGraph', 'size': 12,
         'color': RGBColor(0x7A,0x8B,0xA0), 'align': PP_ALIGN.CENTER}])

# ---------------- 生成 ----------------
slide_cover()
slide_toc()
slide_what_is()
slide_architecture()
slide_paradigms()
slide_tools()
slide_memory()
slide_multiagent()
slide_frameworks()
slide_case()
slide_langgraph()
slide_code()
slide_best()
slide_future()
slide_thanks()

out = '/Users/wxl/Desktop/AI_Agent开发实战指南.pptx'
prs.save(out)
print('OK saved:', out, '| slides:', len(prs.slides._sldIdLst))
